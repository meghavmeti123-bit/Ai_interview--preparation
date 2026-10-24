from __future__ import annotations

import json
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


_LOCK = threading.RLock()
_HEADERS = {
    "Sessions": ["session_id", "created_at", "role", "experience", "interview_type", "focus_areas", "question_count", "user_id", "full_name", "difficulty"],
    "Practice": [
        "question_id", "session_id", "created_at", "focus_area", "question", "sample_answer",
        "options_json", "correct_option", "explanation", "round_number", "role", "role_family",
        "experience_level", "interview_category", "difficulty", "question_type", "format", "concept_id",
        "scenario_id", "bank_question_id", "expected_answer", "evaluation_criteria_json", "tags_json", "option_explanations_json", "source",
        "intentional_retry", "retry_of_question_id", "follow_up_to_question_id",
        "topic", "subtopic", "estimated_time_minutes", "source_reference", "time_limit_seconds",
    ],
    "Reviews": [
        "review_id", "session_id", "question_id", "reviewed_at", "candidate_answer", "score",
        "strengths", "growth_areas", "next_steps", "feedback", "selected_option", "correct_option",
        "is_correct", "explanation", "round_number", "format", "evaluation_json",
        "elapsed_seconds", "timed_out",
    ],
}


def _timestamp() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _open(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, headers in _HEADERS.items():
            workbook.create_sheet(sheet_name).append(headers)
        workbook.save(path)
    workbook = load_workbook(path)
    for sheet_name, headers in _HEADERS.items():
        if sheet_name not in workbook.sheetnames:
            workbook.create_sheet(sheet_name).append(headers)
            continue
        sheet = workbook[sheet_name]
        current_headers = [cell.value for cell in sheet[1]] if sheet.max_row else []
        if not current_headers:
            sheet.append(headers)
        else:
            for header in headers:
                if header not in current_headers:
                    sheet.cell(row=1, column=len(current_headers) + 1, value=header)
                    current_headers.append(header)
    return workbook


def _records(sheet) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, ())
    return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]


def create_session(
    path: Path, setup: dict[str, Any], question_count: int = 5, *,
    user_id: str = "", full_name: str = "",
) -> str:
    with _LOCK:
        workbook = _open(Path(path))
        session_id = str(uuid.uuid4())
        _append_record(workbook["Sessions"], {
            "session_id": session_id,
            "created_at": _timestamp(),
            "role": setup["role"],
            "experience": setup["experience"],
            "interview_type": setup["interview_type"],
            "focus_areas": ", ".join(setup["focus_areas"]),
            "question_count": question_count,
            "user_id": user_id,
            "full_name": full_name,
            "difficulty": setup.get("difficulty", ""),
        })
        workbook.save(path)
        return session_id


def add_practice(
    path: Path,
    session_id: str,
    practice: dict[str, Any],
    round_number: int,
    *,
    intentional_retry: bool = False,
    retry_of_question_id: str = "",
) -> str:
    with _LOCK:
        workbook = _open(Path(path))
        if not intentional_retry:
            from backend.question_engine import is_duplicate

            history = [
                row for row in _records(workbook["Practice"])
                if row["session_id"] == session_id and not bool(row.get("intentional_retry"))
            ]
            if is_duplicate(practice, history):
                raise ValueError("A duplicate interview question cannot be added to this session.")
        question_id = str(uuid.uuid4())
        values = {
            "question_id": question_id,
            "session_id": session_id,
            "created_at": _timestamp(),
            "focus_area": practice["focus_area"],
            "question": practice["question"],
            "sample_answer": practice.get("expected_answer", practice.get("sample_answer", "")),
            "options_json": json.dumps(practice["options"]),
            "correct_option": practice["correct_option"],
            "explanation": practice["explanation"],
            "round_number": round_number,
            "role": practice.get("role", ""),
            "role_family": practice.get("role_family", ""),
            "experience_level": practice.get("experience_level", ""),
            "interview_category": practice.get("interview_category", ""),
            "difficulty": practice.get("difficulty", ""),
            "question_type": practice.get("question_type", ""),
            "format": practice.get("format", "multiple_choice"),
            "concept_id": practice.get("concept_id", ""),
            "scenario_id": practice.get("scenario_id", ""),
            "bank_question_id": practice.get("question_id", ""),
            "expected_answer": practice.get("expected_answer", ""),
            "evaluation_criteria_json": json.dumps(practice.get("evaluation_criteria", [])),
            "tags_json": json.dumps(practice.get("tags", [])),
            "option_explanations_json": json.dumps(practice.get("option_explanations", {})),
            "source": practice.get("source", ""),
            "intentional_retry": intentional_retry,
            "retry_of_question_id": retry_of_question_id,
            "follow_up_to_question_id": practice.get("follow_up_to_question_id", ""),
            "topic": practice.get("topic", practice.get("focus_area", "")),
            "subtopic": practice.get("subtopic", ""),
            "estimated_time_minutes": practice.get("estimated_time_minutes", ""),
            "source_reference": practice.get("source_reference", ""),
            "time_limit_seconds": practice.get("time_limit_seconds", ""),
        }
        _append_record(workbook["Practice"], values)
        workbook.save(path)
        return question_id


def add_review(
    path: Path, session_id: str, question_id: str, answer: str, review: dict[str, Any], round_number: int,
    *, elapsed_seconds: int = 0, timed_out: bool = False,
) -> None:
    with _LOCK:
        workbook = _open(Path(path))
        values = {
            "review_id": str(uuid.uuid4()),
            "session_id": session_id,
            "question_id": question_id,
            "reviewed_at": _timestamp(),
            "candidate_answer": answer,
            "score": int(review.get("score", int(review["is_correct"]))),
            "strengths": json.dumps(["Correct answer selected"] if review["is_correct"] else []),
            "growth_areas": json.dumps([] if review["is_correct"] else [review["focus_area"]]),
            "next_steps": json.dumps([]),
            "feedback": review["explanation"],
            "selected_option": review["selected_option"],
            "correct_option": review["correct_option"],
            "is_correct": review["is_correct"],
            "explanation": review["explanation"],
            "round_number": round_number,
            "format": review.get("format", "multiple_choice"),
            "evaluation_json": json.dumps(review),
            "elapsed_seconds": elapsed_seconds,
            "timed_out": timed_out,
        }
        _append_record(workbook["Reviews"], values)
        workbook.save(path)


def _append_record(sheet, values: dict[str, Any]) -> None:
    headers = [cell.value for cell in sheet[1]]
    sheet.append([values.get(header) for header in headers])


def get_session(path: Path, session_id: str) -> dict[str, Any] | None:
    with _LOCK:
        workbook = _open(Path(path))
        session = next((row for row in _records(workbook["Sessions"]) if row["session_id"] == session_id), None)
        if session is None:
            return None
        practice = [row for row in _records(workbook["Practice"]) if row["session_id"] == session_id]
        reviews = [row for row in _records(workbook["Reviews"]) if row["session_id"] == session_id]
        return {"session": session, "practice": practice, "reviews": reviews}


def get_dashboard_summary(
    path: Path,
    limit: int = 5,
    user_id: str | None = None,
    grade_thresholds: dict[str, int] | None = None,
) -> dict[str, Any]:
    with _LOCK:
        workbook = _open(Path(path))
        sessions = sorted(_records(workbook["Sessions"]), key=lambda row: str(row.get("created_at") or ""), reverse=True)
        practice_rows = _records(workbook["Practice"])
        review_rows = _records(workbook["Reviews"])
        if user_id is not None:
            sessions = [row for row in sessions if str(row.get("user_id") or "") == user_id]
            owned_session_ids = {str(row["session_id"]) for row in sessions}
            practice_rows = [row for row in practice_rows if str(row.get("session_id")) in owned_session_ids]
            review_rows = [row for row in review_rows if str(row.get("session_id")) in owned_session_ids]
        reviews_by_session: dict[str, list[dict[str, Any]]] = {}
        reviews_by_question: dict[tuple[str, str], dict[str, Any]] = {}
        practice_by_session: dict[str, list[dict[str, Any]]] = {}
        for review in review_rows:
            reviews_by_session.setdefault(str(review["session_id"]), []).append(review)
            reviews_by_question[(str(review["session_id"]), str(review["question_id"]))] = review
        for practice in practice_rows:
            if not _as_bool(practice.get("intentional_retry")):
                practice_by_session.setdefault(str(practice["session_id"]), []).append(practice)
        base_question_ids = {
            (str(row["session_id"]), str(row["question_id"]))
            for row in practice_rows
            if not _as_bool(row.get("intentional_retry"))
        }
        base_review_rows = [
            row for row in review_rows
            if (str(row["session_id"]), str(row["question_id"])) in base_question_ids
        ]
        attempted = len(base_review_rows)
        correct = sum(_as_bool(row.get("is_correct")) for row in base_review_rows)
        focus_summary: dict[str, dict[str, int]] = {}
        completed_sessions = 0
        for session in sessions:
            session_id = str(session["session_id"])
            requested_count = int(session.get("question_count") or 0)
            base_practice = [
                row for row in practice_by_session.get(session_id, [])
                if int(row.get("round_number") or 1) == 1
            ]
            reviewed_ids = {
                str(row["question_id"]) for row in base_practice
                if (session_id, str(row["question_id"])) in reviews_by_question
            }
            if requested_count > 0 and len(reviewed_ids) >= requested_count:
                completed_sessions += 1
            for practice in base_practice:
                review = reviews_by_question.get((session_id, str(practice["question_id"])))
                if review is None:
                    continue
                area = str(practice.get("focus_area") or "Other")
                focus_summary.setdefault(area, {"correct": 0, "total": 0})
                focus_summary[area]["total"] += 1
                focus_summary[area]["correct"] += int(_as_bool(review.get("is_correct")))
        recent = []
        completed_attempt_numbers: dict[str, int] = {}
        attempt_number = 0
        for session in reversed(sessions):
            session_id = str(session["session_id"])
            requested_count = int(session.get("question_count") or 0)
            base_practice = [
                row for row in practice_by_session.get(session_id, [])
                if int(row.get("round_number") or 1) == 1
            ]
            reviewed_count = sum(
                (session_id, str(row["question_id"])) in reviews_by_question
                for row in base_practice
            )
            if requested_count > 0 and reviewed_count >= requested_count:
                attempt_number += 1
                completed_attempt_numbers[session_id] = attempt_number
        for session in sessions[:max(0, limit)]:
            session_id = str(session["session_id"])
            session_reviews = [
                row for row in reviews_by_session.get(session_id, [])
                if (session_id, str(row["question_id"])) in base_question_ids
            ]
            session_attempted = len(session_reviews)
            session_correct = sum(_as_bool(row.get("is_correct")) for row in session_reviews)
            session_accuracy = round(100 * session_correct / session_attempted) if session_attempted else None
            recent.append({
                "session_id": session_id,
                "role": session.get("role", "Interview practice"),
                "experience": session.get("experience", ""),
                "created_at": session.get("created_at", ""),
                "full_name": session.get("full_name", ""),
                "requested_question_count": int(session.get("question_count") or 0),
                "attempted_count": session_attempted,
                "correct_count": session_correct,
                "accuracy": session_accuracy,
                "grade": _grade(session_accuracy, grade_thresholds),
                "attempt_number": completed_attempt_numbers.get(session_id),
                "completed": session_id in completed_attempt_numbers,
            })
        return {
            "session_count": len(sessions),
            "question_count": len(base_question_ids),
            "attempted_count": attempted,
            "accuracy": round(100 * correct / attempted) if attempted else None,
            "completed_session_count": completed_sessions,
            "full_name": sessions[0].get("full_name", "") if sessions else "",
            "topic_stats": [
                {"focus_area": area, **counts}
                for area, counts in sorted(
                    focus_summary.items(), key=lambda item: (-item[1]["total"], item[0].casefold())
                )
            ],
            "recent_sessions": recent,
        }


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "true", "yes", "y"}


def _grade(accuracy: int | None, thresholds: dict[str, int] | None = None) -> str | None:
    if accuracy is None:
        return None
    thresholds = thresholds or {"A": 90, "B": 75, "C": 50}
    if accuracy >= int(thresholds["A"]):
        return "A"
    if accuracy >= int(thresholds["B"]):
        return "B"
    if accuracy >= int(thresholds["C"]):
        return "C"
    return "D"